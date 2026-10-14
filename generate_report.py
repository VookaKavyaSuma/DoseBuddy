import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def create_report():
    doc = docx.Document()

    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Styles & Colors
    TEAL = RGBColor(13, 123, 95)       # Primary #0D7B5F
    DARK_SLATE = RGBColor(19, 46, 39)  # Text #132E27
    MUTED = RGBColor(82, 108, 99)      # Secondary #526C63
    AMBER = RGBColor(217, 119, 6)      # Warning #D97706
    BLUE = RGBColor(37, 99, 235)       # Info #2563EB

    # Helper: Set cell background
    def set_cell_background(cell, hex_color):
        shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
        cell._tc.get_or_add_tcPr().append(shading)

    # Helper: Set cell padding
    def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = OxmlElement('w:tcMar')
        for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
            node = OxmlElement(f'w:{m}')
            node.set(qn('w:w'), str(val))
            node.set(qn('w:type'), 'dxa')
            tcMar.append(node)
        tcPr.append(tcMar)

    # Helper: Add Heading 1
    def add_h1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(18)
        run.font.bold = True
        run.font.color.rgb = TEAL
        return p

    # Helper: Add Heading 2
    def add_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(14)
        run.font.bold = True
        run.font.color.rgb = DARK_SLATE
        return p

    # Helper: Add Heading 3
    def add_h3(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(12)
        run.font.bold = True
        run.font.color.rgb = MUTED
        return p

    # Helper: Add Body Paragraph
    def add_p(text, bold_prefix=None):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_bold = p.add_run(bold_prefix)
            r_bold.font.name = 'Calibri'
            r_bold.font.size = Pt(11)
            r_bold.font.bold = True
            r_bold.font.color.rgb = DARK_SLATE
        run = p.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(11)
        run.font.color.rgb = DARK_SLATE
        return p

    # Helper: Add Bullet
    def add_bullet(text, bold_prefix=None):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_bold = p.add_run(bold_prefix)
            r_bold.font.name = 'Calibri'
            r_bold.font.size = Pt(11)
            r_bold.font.bold = True
            r_bold.font.color.rgb = DARK_SLATE
        run = p.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(11)
        run.font.color.rgb = DARK_SLATE
        return p

    # Helper: Callout Box
    def add_callout(title, text, bg_hex="ECFDF5", border_color="10B981"):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.cell(0, 0)
        set_cell_background(cell, bg_hex)
        set_cell_margins(cell, top=140, bottom=140, left=200, right=200)

        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(3)
        r_title = p.add_run(f"{title}\n")
        r_title.font.name = 'Calibri'
        r_title.font.size = Pt(11)
        r_title.font.bold = True
        r_title.font.color.rgb = TEAL

        r_text = p.add_run(text)
        r_text.font.name = 'Calibri'
        r_text.font.size = Pt(10.5)
        r_text.font.italic = True
        r_text.font.color.rgb = DARK_SLATE
        doc.add_paragraph() # Spacer

    # ==========================================
    # COVER / HEADER TITLE
    # ==========================================
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(24)
    p_title.paragraph_format.space_after = Pt(4)
    r_title = p_title.add_run("DoseBuddy")
    r_title.font.name = 'Calibri'
    r_title.font.size = Pt(32)
    r_title.font.bold = True
    r_title.font.color.rgb = TEAL

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(8)
    r_sub = p_sub.add_run("Smart Prescription Reminder & Caregiver Medication Companion")
    r_sub.font.name = 'Calibri'
    r_sub.font.size = Pt(15)
    r_sub.font.bold = True
    r_sub.font.color.rgb = DARK_SLATE

    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_meta.paragraph_format.space_after = Pt(24)
    r_meta = p_meta.add_run("Comprehensive System Functionality, Architecture, and Operational Report\nGenerated: October 2026 • Platform Status: Active & Operational")
    r_meta.font.name = 'Calibri'
    r_meta.font.size = Pt(10)
    r_meta.font.color.rgb = MUTED

    add_callout(
        "⚕️ Regulatory & Medical Safety Disclaimer",
        "DoseBuddy is a medication adherence assistant and reminder schedule companion. It is NOT a licensed healthcare provider, diagnostic engine, or medical device. Extracted details from prescriptions must always be verified by the user, caregiver, or prescribing physician before taking medications.",
        "FEF3C7", "D97706"
    )

    # ==========================================
    # 1. EXECUTIVE SUMMARY & PROBLEM STATEMENT
    # ==========================================
    add_h1("1. Executive Summary & Problem Statement")
    add_p(
        "People, especially older adults and chronic care patients, regularly miss critical medication doses or take pills at the incorrect time of day. This occurs because doctor prescriptions are notoriously difficult to read, medical shorthand (e.g., '1-0-1', 'OD', 'PC') is confusing, and complex multidrug schedules are hard to remember. Furthermore, caregivers often lack real-time visibility into whether their loved ones are staying adherent.",
        bold_prefix="The Healthcare Challenge: "
    )
    add_p(
        "DoseBuddy bridges this gap through a unified full-stack application that transforms doctor-written prescriptions into structured, gentle daily reminders. A user or caregiver scans or uploads a prescription image (including messy cursive doctor slips), our clinical AI engine extracts medicines, dosages, and timings, allows human-in-the-loop verification, and auto-generates a daily schedule partitioned into Morning, Afternoon, and Night slots with proactive 15-minute prior reminders.",
        bold_prefix="The DoseBuddy Solution: "
    )

    # ==========================================
    # 2. SYSTEM ARCHITECTURE & TECH STACK
    # ==========================================
    add_h1("2. System Architecture & Technology Stack")
    add_p(
        "DoseBuddy is built upon an enterprise-grade, decoupled full-stack architecture designed for high reliability, responsiveness, and clinical data integrity:"
    )

    table_tech = doc.add_table(rows=6, cols=3)
    table_tech.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Layer", "Technologies Used", "Key Responsibilities"]
    for i, h in enumerate(headers):
        cell = table_tech.cell(0, i)
        set_cell_background(cell, "0D7B5F")
        set_cell_margins(cell, 120, 120, 150, 150)
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    tech_rows = [
        ("Frontend Web App", "React 19, Vite 8, Vanilla CSS Design System, Web Speech API", "Interactive UI, Daily Schedule slots, Human verification sheet, Speech synthesis, and Browser notifications."),
        ("Backend REST Server", "Java 17, Spring Boot 3.3.5, Spring MVC, Spring Data JPA", "Business logic, prescription scan pipeline, dose action logging, dual caregiver/patient notification routing."),
        ("Security & Auth", "Spring Security 6, Stateless JWT (OAuth2 Resource Server, HMAC-SHA256)", "Role-based access control (Caregiver vs. Patient), session management, and secured REST endpoints."),
        ("Relational Database", "MySQL 8.0 / Hibernate 6.5 with auto DDL updates", "Persistent storage for user profiles, medications, inventory levels, daily dose logs, and notification records."),
        ("AI / Vision & RAG", "Spring AI Gemini Client, Apache PDFBox 3.0, Multimodal Prompts", "Handwriting OCR deciphering, clinical abbreviation translation, and PDF RAG chat assistant.")
    ]

    for row_idx, (col1, col2, col3) in enumerate(tech_rows, start=1):
        r_cells = table_tech.rows[row_idx].cells
        for col_idx, text in enumerate([col1, col2, col3]):
            cell = r_cells[col_idx]
            set_cell_background(cell, "F8FAFC" if row_idx % 2 == 1 else "FFFFFF")
            set_cell_margins(cell, 100, 100, 150, 150)
            p = cell.paragraphs[0]
            r = p.add_run(text)
            r.font.size = Pt(9.5)
            r.font.name = 'Calibri'
            if col_idx == 0:
                r.font.bold = True

    doc.add_paragraph()

    # ==========================================
    # 3. DETAILED FEATURE BREAKDOWN
    # ==========================================
    add_h1("3. Detailed Functional Capabilities")

    add_h2("3.1. AI Prescription Scanner & Messy Handwriting OCR")
    add_bullet("Users can upload image files (JPG, PNG, WEBP), capture photos directly via their device camera, or upload PDF slips.", bold_prefix="Multi-Format Ingestion: ")
    add_bullet("Equipped with specialized clinical prompts and optical filters specifically trained to decode rapid doctor cursive, ink scratches, and difficult medical shorthand.", bold_prefix="Handwriting Deciphering: ")
    add_bullet("Recognizes notations such as '1-0-1' (Morning & Night), '1-1-1' (Thrice daily), '1-0-0' (Morning only), '0-0-1' (Night only), 'OD' (Once daily), 'BD' (Twice daily), 'TDS' (Thrice daily), 'HS' (Bedtime), 'AC' (Before meals), 'PC' (After meals), and 'SOS' (As needed).", bold_prefix="Abbreviation Mapping: ")
    add_bullet("Includes 3 realistic pre-packaged messy doctor prescription presets (Senior Cardiac/Diabetes, Post-Op Antibiotics, and Respiratory) for safe, instant demonstration without exposing real patient health information.", bold_prefix="Safe Sample Presets: ")

    add_h2("3.2. Human-in-the-Loop Review & Confirmation Sheet")
    add_p("Patient safety is fundamental. DoseBuddy strictly requires the user or caregiver to review and confirm all AI-extracted medications before committing them to the daily schedule:")
    add_bullet("Presents the extracted medicines side-by-side with the doctor's name, clinic details, and confidence score.", bold_prefix="Extraction Inspection: ")
    add_bullet("Caregivers can edit medicine names, adjust dosages, toggle Morning/Afternoon/Night time slots, modify meal timing instructions, or add/delete lines.", bold_prefix="Granular Editing: ")
    add_bullet("Upon clicking 'Confirm & Add to Schedule', all validated items are batch-persisted to MySQL and scheduled immediately.", bold_prefix="Batch Synchronization: ")

    add_h2("3.3. Auto-Generated Daily Schedule (Morning, Afternoon, Night)")
    add_bullet("🌅 Morning Slot (08:00 AM): Displays all medications flagged with morning frequency (e.g., Metformin 500mg, Pantoprazole 40mg).", bold_prefix="Morning Regimen: ")
    add_bullet("☀️ Afternoon Slot (01:00 PM): Displays midday doses (e.g., Amoxicillin 500mg midday course).", bold_prefix="Afternoon Regimen: ")
    add_bullet("🌙 Night Slot (08:00 PM): Displays evening and bedtime pills (e.g., Amlodipine 5mg, Atorvastatin 20mg).", bold_prefix="Night Regimen: ")
    add_bullet("Displays pill emojis, exact dosage, meal instructions ('Take after breakfast with water'), and medical purpose tags.", bold_prefix="Elderly-Friendly Visuals: ")

    add_h2("3.4. 15-Minute Prior Reminder Engine")
    add_bullet("The scheduling engine computes the reminder time exactly 15 minutes before medication time (e.g., 07:45 AM for an 08:00 AM morning dose).", bold_prefix="Proactive Notification: ")
    add_bullet("Pops up a focused reminder dialog with an audio chime and Web Speech API voice reading: 'Hello. 15-minute reminder for your medication: Metformin 500mg scheduled at 08:00 AM. Please prepare water and breakfast.'", bold_prefix="In-App Modal & Chime: ")
    add_bullet("Dispatches system desktop notifications via the HTML5 Web Notifications API even if the browser tab is in the background.", bold_prefix="Desktop Push: ")
    add_bullet("A dedicated '🧪 Test 15-Min Prior Reminder Now' button allows users and evaluators to test the 15-minute alert on demand.", bold_prefix="Simulation & Testing: ")

    add_h2("3.5. Dose Tracking with 'Taken' and 'Skip' Reasons")
    add_bullet("Decrements the inventory count, logs a timestamped 'TAKEN' record in the database, and increments adherence metrics.", bold_prefix="Log Dose Taken (-1): ")
    add_bullet("Prompts the user for a clinical skip reason ('Felt unwell / nauseous', 'Physician advised to pause', 'Missed dose window', 'Out of stock', or custom reason).", bold_prefix="Skip Dose Modal: ")
    add_bullet("Records the skip reason in the audit log and immediately alerts the caregiver.", bold_prefix="Caregiver Alerting: ")

    add_h2("3.6. Caregiver On-Demand Editing & Dual Notification Dispatch")
    add_bullet("Caregivers have an '✏️ Edit' button on every card in both the Daily Schedule and Prescription Inventory.", bold_prefix="Instant Editing Access: ")
    add_bullet("Allows editing drug names, strength, '1-0-1' frequencies, exact reminder times, food guidance, and caregiver instructions at any time.", bold_prefix="Comprehensive Regimen Editor: ")
    add_bullet("Saving changes immediately creates two distinct database notifications: one dispatched to the Patient (Robert Jenkins) informing them of the schedule adjustment, and one confirmation dispatched to the Caregiver (Sarah Jenkins).", bold_prefix="Dual Notification Guarantee: ")
    add_bullet("A notification bell in the top navbar displays an unread count badge and opens a slide-out drawer with chronological change summaries.", bold_prefix="Navbar Notification Drawer: ")

    add_h2("3.7. Medication Deletion & Discontinuation with Dual Notification")
    add_bullet("Users and caregivers can delete medications directly from any Daily Schedule slot card (Morning, Afternoon, Night), the Prescription Inventory, or inside the Edit Medication modal.", bold_prefix="Ubiquitous Delete Access: ")
    add_bullet("Requires explicit safety confirmation before removal to prevent accidental discontinuation.", bold_prefix="Accidental Click Protection: ")
    add_bullet("Deleting a medicine immediately dispatches dual notifications: one to the Patient ('Robert Jenkins') and one to the Caregiver ('Sarah Jenkins') documenting that the medicine was removed and upcoming reminders discontinued.", bold_prefix="Dual Notification Guarantee: ")
    add_bullet("All associated upcoming reminders are immediately cleared from the schedule and active pill counts updated in real time.", bold_prefix="Instant Schedule Resynchronization: ")

    add_h2("3.8. Today's Adherence Analytics & Activity Log")
    add_bullet("Displays an adherence score percentage ring (e.g., 75%, 100%) calculated dynamically from today's taken vs. scheduled doses.", bold_prefix="Visual Adherence Ring: ")
    add_bullet("Summary counters for Doses Taken, Doses Skipped, and Remaining Doses for the day.", bold_prefix="Daily KPI Counters: ")
    add_bullet("Interactive timeline listing every dose logged today with exact timestamps, slot names, and recorded skip reasons.", bold_prefix="Audit Timeline: ")
    add_bullet("Displays a motivational streak badge (e.g., '🔥 5-Day Adherence Streak') to encourage patient compliance.", bold_prefix="Gamified Habit Streak: ")

    add_h2("3.9. Informational Drug-Interaction Screening")
    add_bullet("Automatically screens active medications for common co-administration interactions (e.g., Amlodipine + Atorvastatin, Metformin + antihypertensives, Antibiotics + Antacids).", bold_prefix="Automated Screening: ")
    add_bullet("Displays a prominent informational notice with clinical guidance, clearly labeled as 'Informational Only • Not Medical Advice'.", bold_prefix="Safe Labeling: ")

    add_h2("3.10. Inventory Tracking & Low Stock Refill Alerts")
    add_bullet("Every medication tracks total days duration and remaining doses in the pillbox.", bold_prefix="Pill Count Tracking: ")
    add_bullet("Automatically switches the inventory bar to warning state and displays '⚠️ Low Stock Alert: Refill Needed Soon!' whenever remaining doses drop to 5 or fewer.", bold_prefix="Automatic Refill Warning: ")

    add_h2("3.11. Multi-Language Accessibility")
    add_bullet("Includes a top navbar language dropdown supporting English (🇬🇧), Spanish (🇪🇸), Hindi (🇮🇳), and French (🇫🇷).", bold_prefix="Language Switcher: ")
    add_bullet("Translates core UI headers, time slots ('🌅 Morning Slot' -> '🌅 सुबह की खुराक' -> '🌅 Turno de la Mañana'), and dosage instructions for senior accessibility.", bold_prefix="Localized Schedules: ")

    add_h2("3.12. AI Conversational Assistant & Document RAG")
    add_bullet("A floating pill icon widget ('💊 DoseBuddy AI') provides conversational answers about food warnings, side effects, and missed dose guidance.", bold_prefix="Floating Health Companion: ")
    add_bullet("Supports uploading medical PDF reports, indexing them in memory via chunking, and grounding AI responses in the user's specific health documents.", bold_prefix="Prescription PDF RAG: ")

    add_h2("3.13. Universal Responsive Mobile & Tablet Layout")
    add_bullet("Specially engineered for seniors and caregivers using smartphones, iPads/tablets, and desktop screens with seamless viewport adaptation.", bold_prefix="Fluid Viewport Adaptation: ")
    add_bullet("Action buttons automatically convert to large, thumb-friendly touch targets (min 44px height) on mobile devices to prevent mis-clicks by older adults.", bold_prefix="Senior-Friendly Touch Targets: ")
    add_bullet("Slide-out notification drawer and modal dialogs expand gracefully to full-width mobile viewports with native app ergonomics.", bold_prefix="Native-Feel Mobile UX: ")

    # ==========================================
    # 4. DATABASE SCHEMA
    # ==========================================
    add_h1("4. Database Architecture & Schema")
    add_p("The MySQL database ('dosebuddy_db') contains five core relational tables:")

    table_schema = doc.add_table(rows=6, cols=3)
    table_schema.alignment = WD_TABLE_ALIGNMENT.CENTER
    s_headers = ["Table Name", "Key Columns", "Purpose"]
    for i, h in enumerate(s_headers):
        cell = table_schema.cell(0, i)
        set_cell_background(cell, "0D7B5F")
        set_cell_margins(cell, 120, 120, 150, 150)
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    schema_rows = [
        ("users", "id, email, username, full_name, password, role", "Stores registered caregivers and patients with hashed BCrypt passwords and roles."),
        ("medications", "id, name, dosage, frequency, morning_dose, afternoon_dose, night_dose, remaining_doses, category", "Stores active and past prescriptions, scheduling flags, refill counts, and caregiver notes."),
        ("dose_logs", "id, medication_id, medication_name, time_slot, status, reason, log_date, logged_at", "Chronological audit log tracking each dose taken or skipped, time of action, and patient skip reason."),
        ("notification_records", "id, recipient_role, recipient_name, medication_name, title, message, change_summary, is_read", "Maintains dual dispatch notifications for both caregiver and patient when medicines change."),
        ("chat_messages", "id, conversation_id, role, content, timestamp", "Persistent history of conversational interactions between the user and DoseBuddy AI.")
    ]

    for row_idx, (col1, col2, col3) in enumerate(schema_rows, start=1):
        r_cells = table_schema.rows[row_idx].cells
        for col_idx, text in enumerate([col1, col2, col3]):
            cell = r_cells[col_idx]
            set_cell_background(cell, "F8FAFC" if row_idx % 2 == 1 else "FFFFFF")
            set_cell_margins(cell, 100, 100, 150, 150)
            p = cell.paragraphs[0]
            r = p.add_run(text)
            r.font.size = Pt(9.5)
            r.font.name = 'Calibri'
            if col_idx == 0:
                r.font.bold = True

    doc.add_paragraph()

    # ==========================================
    # 5. REST API SPECIFICATIONS
    # ==========================================
    add_h1("5. REST API Specifications")
    add_bullet("POST /api/prescriptions/scan: Ingests prescription image/PDF or preset ID; returns structured medications, dosages, 1-0-1 frequencies, and doctor metadata.", bold_prefix="Prescription Scanner: ")
    add_bullet("GET /api/prescriptions/samples: Retrieves sample messy prescription presets for demonstration.", bold_prefix="Sample Presets: ")
    add_bullet("GET /api/medications: Retrieves active or past medications ordered by creation date.", bold_prefix="Fetch Medications: ")
    add_bullet("POST /api/medications/batch: Saves confirmed batch of medications from the scanner review sheet.", bold_prefix="Batch Creation: ")
    add_bullet("PUT /api/medications/{id}: Caregiver updates medication; generates diff and dispatches dual notifications to patient and caregiver.", bold_prefix="Edit Medication: ")
    add_bullet("DELETE /api/medications/{id}: Deletes medication from schedule and MySQL inventory; dispatches dual notifications to both patient and caregiver and cancels upcoming reminders.", bold_prefix="Delete Medication: ")
    add_bullet("POST /api/medications/{id}/take: Decrements dose count, logs 'TAKEN' event, and updates adherence.", bold_prefix="Log Taken Dose: ")
    add_bullet("POST /api/medications/{id}/skip: Records dose skip with patient reason and notifies caregiver.", bold_prefix="Log Skipped Dose: ")
    add_bullet("GET /api/medications/logs: Returns today's dose logs for adherence calculations and audit timeline.", bold_prefix="Fetch Dose Logs: ")
    add_bullet("GET /api/medications/notifications: Fetches recent notifications filtered by recipient role (PATIENT/CAREGIVER).", bold_prefix="Fetch Notifications: ")
    add_bullet("POST /api/chat: Conversational medication Q&A with RAG document grounding.", bold_prefix="AI Chatbot: ")

    # ==========================================
    # 6. SUMMARY & CONCLUSION
    # ==========================================
    add_h1("6. Summary & Conclusion")
    add_p(
        "DoseBuddy delivers an end-to-end, empathetic, and clinically informed software solution for prescription adherence. By automating the transition from confusing, messy handwritten prescriptions into clear, time-slotted daily reminders, DoseBuddy substantially reduces the likelihood of missed or mistimed doses among senior patients. With real-time dual notifications, caregivers remain continuously informed and empowered, creating a collaborative, transparent care ecosystem."
    )

    return doc

if __name__ == '__main__':
    doc = create_report()
    
    # Save to workspace root
    workspace_path = r"c:\Users\Dell\Desktop\DoseBuddy\DoseBuddy_System_Comprehensive_Report.docx"
    try:
        doc.save(workspace_path)
        print(f"Saved workspace report to: {workspace_path}")
    except PermissionError:
        workspace_path = r"c:\Users\Dell\Desktop\DoseBuddy\DoseBuddy_System_Comprehensive_Report_v2.docx"
        doc.save(workspace_path)
        print(f"Original file was open in Word. Saved updated report to: {workspace_path}")

    # Save to artifacts directory
    artifacts_dir = r"C:\Users\Dell\.gemini\antigravity-ide\brain\86bf23fd-e463-4895-91e1-b378536b2620"
    artifact_path = os.path.join(artifacts_dir, "DoseBuddy_System_Comprehensive_Report.docx")
    try:
        doc.save(artifact_path)
        print(f"Saved artifact report to: {artifact_path}")
    except Exception as e:
        print(f"Artifact save note: {e}")
